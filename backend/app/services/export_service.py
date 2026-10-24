"""Record export service — CSV and Excel."""

import csv
import io
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models.crime import Incident

_COLUMNS = [
    ("crime_id", "Crime ID"),
    ("case_number", "FIR No"),
    ("primary_type", "Crime Type"),
    ("description", "Description"),
    ("location_description", "Location Type"),
    ("block", "Address"),
    ("latitude", "Latitude"),
    ("longitude", "Longitude"),
    ("police_station", "Police Station"),
    ("zone_name", "Zone"),
    ("occurred_at", "Date/Time"),
    ("risk_score", "Risk Score"),
    ("risk_level", "Risk Level"),
    ("status", "Status"),
    ("arrest", "Arrest"),
]


def _row(inc: Incident) -> list:
    return [
        inc.crime_id,
        inc.case_number or "",
        inc.primary_type,
        inc.description or "",
        inc.location_description or "",
        inc.block or "",
        inc.latitude,
        inc.longitude,
        inc.police_station or "",
        inc.zone_name or "",
        inc.occurred_at.strftime("%Y-%m-%d %H:%M") if inc.occurred_at else "",
        inc.risk_score,
        inc.risk_level,
        inc.status,
        "Yes" if inc.arrest else "No",
    ]


def to_csv(incidents: Sequence[Incident]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([h for _, h in _COLUMNS])
    for inc in incidents:
        writer.writerow(_row(inc))
    return buf.getvalue().encode("utf-8")


def to_excel(incidents: Sequence[Incident]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Crime Records"

    header_fill = PatternFill("solid", fgColor="1E3A8A")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append([h for _, h in _COLUMNS])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for inc in incidents:
        ws.append(_row(inc))

    for i, (_, header) in enumerate(_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(
            14, len(header) + 2
        )

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
